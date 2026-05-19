import csv
import io
from datetime import datetime

from django.http import HttpResponse


TASK_COLUMNS = [
    'ID', 'Title', 'Status', 'Priority', 'Assignee', 'Assignees',
    'Sprint', 'Milestone', 'Project',
    'Due Date', 'Start Date',
    'Estimated Hours', 'Actual Hours',
    'Tags', 'Is Blocked', 'Created At',
]


def _task_row(task):
    assignee = task.assignee.username if task.assignee else ''
    assignees = ', '.join(u.username for u in task.assignees.all())
    sprint = task.sprint.name if task.sprint else ''
    milestone = task.sprint.milestone.name if task.sprint and task.sprint.milestone else ''
    project = task.sprint.milestone.project.name if task.sprint and task.sprint.milestone and task.sprint.milestone.project else ''
    tags = ', '.join(t.name for t in task.tags.all())
    return [
        task.id,
        task.title,
        task.status,
        task.priority,
        assignee,
        assignees,
        sprint,
        milestone,
        project,
        str(task.due_date) if task.due_date else '',
        str(task.start_date) if task.start_date else '',
        str(task.estimated_hours) if task.estimated_hours is not None else '',
        str(task.actual_hours),
        tags,
        'Yes' if getattr(task, '_is_blocked', False) else '',
        task.created_at.strftime('%Y-%m-%d %H:%M') if task.created_at else '',
    ]


def generate_csv_response(tasks, filename='tasks'):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
    writer = csv.writer(response)
    writer.writerow(TASK_COLUMNS)
    for task in tasks:
        writer.writerow(_task_row(task))
    return response


def generate_xlsx_response(tasks, filename='tasks'):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Tasks'

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')

    for col_idx, col_name in enumerate(TASK_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    for row_idx, task in enumerate(tasks, start=2):
        for col_idx, value in enumerate(_task_row(task), start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    # Auto-fit column widths (approximate)
    for col_idx, col_name in enumerate(TASK_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(len(col_name) + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    return response
